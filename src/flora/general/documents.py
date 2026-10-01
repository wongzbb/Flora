# SPDX-License-Identifier: Apache-2.0
"""Bounded document extraction, deterministic table operations and report publication."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import subprocess
import uuid
import zipfile
from decimal import ROUND_HALF_EVEN, Decimal, InvalidOperation, localcontext
from pathlib import Path, PurePosixPath

from flora.integrations.workspace import WorkspaceTools, _publication_guard
from flora.support.errors import InterruptedEffect, ValidationError

from ._document_process import run_worker

MAX_DOCUMENT = 8 * 1024 * 1024
MAX_TEXT = 3 * 1024 * 1024
MAX_ROWS = 20000


class DocumentWorkspace(WorkspaceTools):
    def _write(self, parts, content, *, expected_sha256=None, create=False, transform=None):
        result = super()._write(
            parts, content, expected_sha256=expected_sha256, create=create, transform=transform
        )
        callback = getattr(self, "_published_callback", None)
        if callback is not None:
            try:
                callback(result)
            except Exception as exc:
                raise InterruptedEffect(
                    "File was published but its artifact receipt could not be persisted"
                ) from exc
        return result

    def create_file(self, path: str, content: str) -> dict:
        """Create a new UTF-8 file atomically; never overwrite an existing destination.

        Parent directories must exist. Returns the actual publication path and sha256.
        An existing destination raises FileExistsError, with no overwrite/publication.
        """
        try:
            return self.write_file(path, content, create=True)
        except ValidationError as exc:
            # Preserve the actual OS conflict type for this new explicit API.
            # Legacy write_file keeps its original ValidationError contract; no
            # error-message matching, preflight race or additional write is used.
            if isinstance(exc.__cause__, FileExistsError):
                raise FileExistsError(
                    "Destination already exists; no file was overwritten"
                ) from exc
            raise

    def update_file(self, path: str, content: str, expected_sha256: str) -> dict:
        """Replace full UTF-8 content using the full sha256 from an actual read_file.

        Refuses stale hashes. Returns the actual publication path and sha256.
        """
        return self.write_file(path, content, expected_sha256=expected_sha256)

    def append_lines(self, path: str, lines: list[str], expected_sha256: str) -> dict:
        """Append lines atomically, preserving all existing bytes and checking the observed hash.

        Prefer for adding lines: inserts a separator only if missing, then terminates
        each new line. Uses the existing final newline style (otherwise LF). Supply
        1..10000 lines without embedded CR/LF, and the full sha256 from read_file.
        Returns the actual publication path and sha256; never creates a missing file.
        """
        if (
            type(lines) is not list
            or not 1 <= len(lines) <= 10000
            or any(type(line) is not str or "\r" in line or "\n" in line for line in lines)
        ):
            raise ValidationError("lines must contain 1..10000 strings without CR/LF")

        def transform(previous):
            ending = (
                "\r\n" if previous.endswith("\r\n") else ("\r" if previous.endswith("\r") else "\n")
            )
            separator = "" if not previous or previous.endswith(("\n", "\r")) else ending
            return previous + separator + ending.join(lines) + ending

        return self._write(
            self._parts(path), None, expected_sha256=expected_sha256, transform=transform
        )

    def read_bytes(self, path):
        parts = self._parts(path)
        with self._directory(parts[:-1]) as parent:
            raw, _ = self._read(parent, parts[-1], MAX_DOCUMENT)
        if len(raw) > MAX_DOCUMENT:
            raise ValidationError("Document exceeds the 8 MiB input limit")
        return raw

    def create_bytes(self, path, data):
        """Publish a new attachment atomically; never overwrite an existing file."""
        if not isinstance(data, bytes) or len(data) > MAX_DOCUMENT:
            raise ValidationError("Binary artifact exceeds the 8 MiB limit")
        parts = self._parts(path)
        with _publication_guard() as publication, self._directory(
            parts[:-1], write_lock=True
        ) as parent:
            temporary = ".flora-write-" + uuid.uuid4().hex
            fd = os.open(
                temporary,
                os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW,
                0o600,
                dir_fd=parent,
            )
            try:
                with os.fdopen(fd, "wb") as stream:
                    stream.write(data)
                    stream.flush()
                    os.fsync(stream.fileno())
                os.link(
                    temporary,
                    parts[-1],
                    src_dir_fd=parent,
                    dst_dir_fd=parent,
                    follow_symlinks=False,
                )
                publication["published"] = True
                os.fsync(parent)
            except FileExistsError:
                raise ValidationError("Artifact already exists; select a new path") from None
            finally:
                os.unlink(temporary, dir_fd=parent)
        return {
            "path": "/".join(parts),
            "sha256": hashlib.sha256(data).hexdigest(),
            "size_bytes": len(data),
        }

    def make_directory(self, path: str) -> dict:
        """Create a workspace directory and its parents without following symbolic links."""
        parts = self._parts(path, directory=True)
        for i, part in enumerate(parts):
            with self._directory(parts[:i], write_lock=True) as parent:
                try:
                    os.mkdir(part, mode=0o700, dir_fd=parent)
                    os.fsync(parent)
                except FileExistsError:
                    pass
            with self._directory(parts[: i + 1]):
                pass
        return {"path": "/".join(parts), "exists": True}


def check_archive(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(x.file_size for x in entries) > 32 * 1024 * 1024:
            raise ValidationError("Document archive exceeds decompressed limits")
        for entry in entries:
            if entry.flag_bits & 1 or entry.file_size > 12 * 1024 * 1024:
                raise ValidationError("Encrypted or oversized archive member")
            if entry.file_size > max(1, entry.compress_size) * 500:
                raise ValidationError("Document archive compression ratio exceeds limit")


def _checked_text(text):
    if len(text.encode("utf-8")) > MAX_TEXT:
        raise ValidationError("Extracted text exceeds the 3 MiB limit; split the document first")
    return text


def _parse_worker(raw, suffix, operation):
    if not isinstance(raw, bytes) or len(raw) > MAX_DOCUMENT:
        raise ValidationError("Document exceeds the 8 MiB input limit")
    # Parsing untrusted document containers has separate memory/CPU/wall bounds.
    # The worker receives neither model credentials nor inherited service tokens.
    environment = {
        key: os.environ[key]
        for key in ("PATH", "LANG", "LC_ALL", "SYSTEMROOT")
        if key in os.environ
    }
    environment["PYTHONPATH"] = str(Path(__file__).resolve().parents[2])
    payload = json.dumps({"suffix": suffix, "operation": operation}).encode() + b"\n" + raw
    try:
        result = run_worker(payload, environment)
    except subprocess.TimeoutExpired:
        raise ValidationError("Document parsing exceeded the 30-second deadline") from None
    if result.returncode != 0 or len(result.stdout) > 12 * 1024 * 1024:
        raise ValidationError("Document parser exceeded resource limits or stopped unexpectedly")
    try:
        response = json.loads(result.stdout)
    except (ValueError, UnicodeError):
        raise ValidationError("Document parser returned an invalid result") from None
    if "error" in response:
        raise ValidationError(response["error"])
    return response["value"]


def extract_document(raw, suffix):
    return _parse_worker(raw, suffix, "extract")


def read_tables(raw, suffix):
    return _parse_worker(raw, suffix, "tables")


def _extract_document(raw, suffix):
    if len(raw) > MAX_DOCUMENT:
        raise ValidationError("Document exceeds the 8 MiB input limit")
    suffix = suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValidationError("Encrypted PDFs must be decrypted before attachment")
            if len(reader.pages) > 300:
                raise ValidationError("PDF exceeds 300 pages; split it first")
            parts = []
            for i, page in enumerate(reader.pages):
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > 8 * 1024 * 1024:
                    raise ValidationError("PDF page content exceeds extraction limits")
                parts.append(f"[Page {i + 1}]\n" + (page.extract_text() or ""))
                _checked_text("\n".join(parts))
            text = "\n\n".join(parts)
            return _checked_text(text), {
                "pages": len(parts),
                "ocr_performed": False,
                "text_characters": sum(len(p.split("\n", 1)[-1]) for p in parts),
            }
        if suffix == ".docx":
            from docx import Document

            check_archive(raw)
            document = Document(io.BytesIO(raw))
            parts = [p.text for p in document.paragraphs]
            for i, table in enumerate(document.tables):
                parts.append(f"[Table {i + 1}]")
                parts.extend("\t".join(c.text for c in row.cells) for row in table.rows)
            return _checked_text("\n".join(parts)), {
                "tables": len(document.tables),
                "images_extracted": False,
            }
        if suffix in {".xlsx", ".csv"}:
            sheets = _read_tables(raw, suffix)
            text = "\n\n".join(
                f"[Sheet {name}]\n" + json.dumps(rows, ensure_ascii=False, default=str)
                for name, rows in sheets.items()
            )
            return _checked_text(text), {
                "sheets": {name: len(rows) for name, rows in sheets.items()},
                "formulas_recalculated": False,
            }
        if suffix not in {
            ".txt",
            ".md",
            ".json",
            ".jsonl",
            ".xml",
            ".html",
            ".log",
            ".yaml",
            ".yml",
            ".rst",
        }:
            raise ValidationError(
                "Supported documents: PDF, DOCX, XLSX, CSV and UTF-8 text formats"
            )
        text = raw.decode("utf-8-sig")
        if "\0" in text:
            raise ValidationError("Binary content is not a text document")
        return _checked_text(text), {"encoding": "utf-8"}
    except ImportError:
        raise ValidationError(
            "Install document dependencies with: pip install 'flora-lang[general]'"
        ) from None
    except (ValueError, UnicodeError, zipfile.BadZipFile, KeyError) as exc:
        raise ValidationError("Document could not be parsed: " + type(exc).__name__) from exc


def _read_tables(raw, suffix):
    if suffix == ".csv":
        reader = csv.reader(io.StringIO(raw.decode("utf-8-sig")))
        rows = []
        for row in reader:
            if len(rows) >= MAX_ROWS or len(row) > 256:
                raise ValidationError("Table exceeds 20,000 rows or 256 columns")
            rows.append(row)
        return {"Sheet1": rows}
    if suffix != ".xlsx":
        raise ValidationError("Table input must be CSV or XLSX")
    from openpyxl import load_workbook

    check_archive(raw)
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        if len(book.worksheets) > 32:
            raise ValidationError("Workbook exceeds 32 sheets")
        result, total = {}, 0
        for sheet in book.worksheets:
            if (
                sheet.max_column
                and sheet.max_column > 256
                or sheet.max_row
                and sheet.max_row > MAX_ROWS
            ):
                raise ValidationError("Worksheet dimensions exceed limits")
            rows = []
            for row in sheet.iter_rows(values_only=True):
                total += 1
                if total > MAX_ROWS or len(row) > 256:
                    raise ValidationError("Workbook exceeds 20,000 rows or 256 columns")
                rows.append(list(row))
            result[sheet.title] = rows
        return result
    finally:
        book.close()


class DocumentTools:
    def __init__(self, files, store, task_key):
        self.files, self.store, self.task_key = files, store, task_key

    def read_document(self, path: str) -> dict:
        """Extract PDF/DOCX/XLSX/CSV/text into a saved source. Use read_source to page through it."""
        raw = self.files.read_bytes(path)
        text, details = extract_document(raw, PurePosixPath(path).suffix)
        result = self.store.record(origin="workspace:" + path, title=path, text=text, raw=raw)
        return {**result, "details": details}

    def table_query(
        self,
        path: str,
        sheet: str = "",
        filters: list | None = None,
        group_by: list | None = None,
        metrics: list | None = None,
        offset: int = 0,
        limit: int = 100,
    ) -> dict:
        """Query CSV/XLSX. First row is unique column names. Filters: {column,op,value}, op eq/ne/gt/ge/lt/le/contains.
        Metrics: {op:count|sum|mean|min|max,column,as}; count may omit column. Numeric arithmetic uses Decimal.
        With metrics, group_by names group columns. Without metrics, page matching rows. No formula recalculation.
        """
        if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 500:
            raise ValidationError("Table offset must be >= 0 and limit must be between 1 and 500")
        raw = self.files.read_bytes(path)
        tables = read_tables(raw, PurePosixPath(path).suffix.lower())
        if not tables:
            raise ValidationError("No sheets")
        selected = sheet or next(iter(tables))
        if selected not in tables:
            raise ValidationError("Unknown sheet: " + selected)
        cells = tables[selected]
        if not cells:
            raise ValidationError("Table has no header row")
        names = [str(x) if x is not None else "" for x in cells[0]]
        if any(not x for x in names) or len(names) != len(set(names)):
            raise ValidationError("Column names must be nonempty and unique")
        if any(len(r) != len(names) for r in cells[1:]):
            raise ValidationError("Table has inconsistent row widths")
        rows = [dict(zip(names, row)) for row in cells[1:]]
        filters, groups, metrics = filters or [], group_by or [], metrics or []
        if (
            not all(isinstance(x, list) for x in (filters, groups, metrics))
            or len(filters) > 20
            or len(metrics) > 20
        ):
            raise ValidationError("Invalid table query specification")
        if any(x not in names for x in groups) or len(groups) != len(set(groups)):
            raise ValidationError("Unknown or duplicate group column")

        def decimal(value):
            try:
                if len(str(value)) > 128:
                    raise InvalidOperation
                number = Decimal(str(value))
                if not number.is_finite() or not -256 <= number.as_tuple().exponent <= 256:
                    raise InvalidOperation
                return number
            except InvalidOperation:
                raise ValidationError(
                    "Numeric operation encountered missing or nonnumeric data"
                ) from None

        for f in filters:
            if (
                not isinstance(f, dict)
                or set(f) != {"column", "op", "value"}
                or f["column"] not in names
            ):
                raise ValidationError("Invalid filter")
            op, value = f["op"], f["value"]
            if op not in {"eq", "ne", "gt", "ge", "lt", "le", "contains"}:
                raise ValidationError("Unknown filter operator")

            def keep(row):
                v = row[f["column"]]
                if op == "eq":
                    return str(v) == str(value)
                if op == "ne":
                    return str(v) != str(value)
                if op == "contains":
                    return str(value) in str(v)
                a, b = decimal(v), decimal(value)
                return {"gt": a > b, "ge": a >= b, "lt": a < b, "le": a <= b}[op]

            rows = [row for row in rows if keep(row)]
        matched = len(rows)
        if metrics:
            aliases = []
            for metric in metrics:
                if not isinstance(metric, dict) or set(metric) - {"column", "op", "as"}:
                    raise ValidationError("Invalid metric")
                if metric.get("op") not in {"count", "sum", "mean", "min", "max"}:
                    raise ValidationError("Unknown metric operator")
                if metric["op"] != "count" and metric.get("column") not in names:
                    raise ValidationError("Metric requires a known column")
                alias = metric.get("as", metric["op"] + "_" + metric.get("column", "rows"))
                if not isinstance(alias, str) or not alias or alias in groups + aliases:
                    raise ValidationError("Metric aliases must be nonempty and unique")
                aliases.append(alias)
            buckets = {}
            for row in rows:
                key = tuple(str(row[x]) for x in groups)
                buckets.setdefault(key, []).append(row)
            if not groups and not buckets:
                buckets[()] = []
            output = []
            for key, values in buckets.items():
                item = dict(zip(groups, key))
                for metric, alias in zip(metrics, aliases):
                    op = metric["op"]
                    nums = (
                        [decimal(row[metric["column"]]) for row in values] if op != "count" else []
                    )
                    if op == "count":
                        result = len(values)
                    elif op == "sum":
                        with localcontext() as context:
                            context.prec = 1024
                            result = str(sum(nums, Decimal(0)))
                    elif not nums:
                        result = None
                    elif op == "mean":
                        with localcontext() as context:
                            context.prec = 1024
                            total = sum(nums, Decimal(0))
                            context.prec = 34
                            context.rounding = ROUND_HALF_EVEN
                            result = str(total / len(nums))
                    elif op == "min":
                        result = str(min(nums))
                    else:
                        result = str(max(nums))
                    item[alias] = result
                output.append(item)
        else:
            if groups:
                raise ValidationError("group_by requires metrics")
            output = rows
        query = {"sheet": selected, "filters": filters, "group_by": groups, "metrics": metrics}
        text = json.dumps(
            {
                "query": query,
                "input_sha256": hashlib.sha256(raw).hexdigest(),
                "matched_rows": matched,
                "numeric_policy": "Bounded decimal inputs; exact sums; means rounded to 34 significant digits (half even)",
                "rows": output,
            },
            ensure_ascii=False,
            default=str,
        )
        saved = self.store.record(
            origin="table:"
            + path
            + ":"
            + hashlib.sha256(json.dumps(query, sort_keys=True).encode()).hexdigest()[:16],
            title="Table query: " + path,
            text=_checked_text(text),
        )
        selected_rows = []
        page_bytes = 0
        for row in output[offset : offset + limit]:
            cost = len(json.dumps(row, ensure_ascii=False, default=str).encode())
            if page_bytes + cost > 60000:
                break
            selected_rows.append(row)
            page_bytes += cost
        rows_in_source = bool(offset < len(output) and not selected_rows)
        next_offset = offset + len(selected_rows)
        return {
            "source_id": saved["source_id"],
            "input_sha256": hashlib.sha256(raw).hexdigest(),
            "rows": json.loads(json.dumps(selected_rows, default=str)),
            "rows_in_source": rows_in_source,
            "matched_rows": matched,
            "total_result_rows": len(output),
            "next_offset": next_offset
            if next_offset < len(output) and not rows_in_source
            else None,
            "formulas_recalculated": False,
            "numeric_policy": "Bounded decimal inputs; exact sums; means rounded to 34 significant digits (half even)",
        }

    def write_report(
        self, path: str, content: str, source_ids: list[str], expected_sha256: str | None = None
    ) -> dict:
        """Write a Markdown report with [src-000001] citations and an appended source ledger.
        Every cited source must exist and source_ids must exactly match the citations. Create by default;
        replace only with an observed expected_sha256. This checks references, not claim correctness.
        """
        if not isinstance(content, str) or not content.strip() or len(content.encode()) > 1048576:
            raise ValidationError("Report content must be nonempty and at most 1 MiB")
        if (
            not isinstance(source_ids, list)
            or len(source_ids) > 1000
            or any(not isinstance(x, str) for x in source_ids)
        ):
            raise ValidationError("source_ids must be a list of source IDs")
        cited = set(re.findall(r"\[(src-[^\]\s]+)\]", content))
        if cited != set(source_ids) or len(source_ids) != len(set(source_ids)):
            raise ValidationError("source_ids must exactly match the unique inline citations")
        references = []
        for ident in source_ids:
            source = self.store.read_source(ident, limit=1)
            references.append(
                f"- [{ident}] {json.dumps(source['title'], ensure_ascii=False)} — "
                f"{json.dumps(source['origin'], ensure_ascii=False)}; SHA-256: {source['sha256']}"
            )
        text = (
            content.rstrip()
            + ("\n\n## Sources\n\n" + "\n".join(references) if references else "")
            + "\n"
        )
        result = self.files.write_file(
            path, text, expected_sha256=expected_sha256, create=expected_sha256 is None
        )
        try:
            self.store.record_artifact(result["path"], result["sha256"], source_ids, self.task_key())
        except Exception as exc:
            raise InterruptedEffect(
                "Report was published but its artifact receipt could not be persisted; do not retry"
            ) from exc
        return {
            **result,
            "source_ids": source_ids,
            "reference_check": "passed",
            "claims_verified": False,
        }

    def export_document(self, source_path: str, output_path: str) -> dict:
        """Export a UTF-8 text/Markdown report to DOCX or PDF, or CSV to XLSX. Output must be new.
        Text formatting is line/heading based; arbitrary Markdown HTML, scripts and embedded media are not executed.
        """
        raw = self.files.read_bytes(source_path)
        suffix = PurePosixPath(output_path).suffix.lower()
        stream = io.BytesIO()
        if suffix == ".xlsx":
            from openpyxl import Workbook

            if PurePosixPath(source_path).suffix.lower() != ".csv":
                raise ValidationError("XLSX export requires a CSV source")
            book = Workbook()
            for row in read_tables(raw, ".csv")["Sheet1"]:
                book.active.append(row)
                for cell in book.active[book.active.max_row]:
                    cell.data_type = "s"  # Untrusted strings never become spreadsheet formulas.
            book.save(stream)
            book.close()
        elif suffix == ".docx":
            from docx import Document

            document = Document()
            for line in raw.decode("utf-8-sig").splitlines():
                m = re.match(r"^(#{1,6})\s+(.*)", line)
                if m:
                    document.add_heading(m[2], level=len(m[1]))
                else:
                    document.add_paragraph(line)
            document.save(stream)
        elif suffix == ".pdf":
            from xml.sax.saxutils import escape

            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            style = getSampleStyleSheet()["BodyText"]
            style.fontName, style.leading, style.wordWrap = "STSong-Light", 15, "CJK"
            parts = []
            for line in raw.decode("utf-8-sig").splitlines():
                parts.append(Paragraph(escape(line) or " ", style))
                parts.append(Spacer(1, 3))
            SimpleDocTemplate(stream).build(parts)
        else:
            raise ValidationError("Export extension must be .docx, .pdf or .xlsx")
        sources = []
        for item in self.store.artifacts():
            if item["path"] == source_path and item["sha256"] == hashlib.sha256(raw).hexdigest():
                sources = item["sources"]
        result = self.files.create_bytes(output_path, stream.getvalue())
        try:
            self.store.record_artifact(
                result["path"], result["sha256"], sources, self.task_key(), kind="export"
            )
        except Exception as exc:
            raise InterruptedEffect(
                "Document was exported but its artifact receipt could not be persisted; do not retry"
            ) from exc
        return result
